"""INT128, NUMERIC(38), DECFLOAT and TIME/TIMESTAMP WITH TIME ZONE through the driver patches
(fbadmin/fb4types.py, fbadmin/driver_patches.py) and the API. Firebird 4+ only."""

import datetime as dt
import decimal
import io
import json
import time
import uuid
import zoneinfo

import firebirdsql
import pytest

import fbadmin  # noqa: F401 - applies the driver patches
from fbadmin import fb4types
from fbadmin.serialize import fb_temporal_text, to_json

from .conftest import DB_DIR, HOST, PASSWORD, PORT, USER, requires_server

D = decimal.Decimal
NY = zoneinfo.ZoneInfo("America/New_York")
BUC = zoneinfo.ZoneInfo("Europe/Bucharest")
P0530 = dt.timezone(dt.timedelta(hours=5, minutes=30))
M0300 = dt.timezone(dt.timedelta(hours=-3))

TABLE = ("create table t (id int not null primary key, i int128, n numeric(38,4), d decimal(38,0), "
         "f16 decfloat(16), f34 decfloat(34), tt time with time zone, ts timestamp with time zone)")
# id, INT128, NUMERIC(38,4), DECIMAL(38,0), DECFLOAT(16), DECFLOAT(34), TIME TZ, TIMESTAMP TZ (SQL text)
ROWS = [
    (1, "170141183460469231731687303715884105727", "-1234567890123456789012345678901234.5678",
     "12345678901234567890123456789012345678", "1.234567890123456E+300", "-1.234567890123456789012345678901234E-6000",
     "10:11:12.3456 America/New_York", "2024-07-01 10:11:12.3456 America/New_York"),
    (2, "-170141183460469231731687303715884105728", "-0.0001", "-98765432109876543210", "-0", "NaN",
     "23:59:59 Europe/Bucharest", "2024-01-15 23:59:59 Europe/Bucharest"),
    (3, "18446744073709551616", "-1.5", "-1", "Infinity", "-Infinity", "01:00 +05:30", "2024-03-10 01:00 +05:30"),
    (4, "-18446744073709551617", "12.34", "9223372036854775808", "1E+384", "1E-6143", "00:00 -03:00",
     "0001-01-01 00:00 -03:00"),
    (5, "-1", "0", "0", "123.45", "0.1", "12:00 UTC", "9999-12-31 23:59:59.9999 UTC"),
    (6, "9223372036854775807", "99999999999999999999999999999999.9999", "-2", "9999999999999999",
     "-8.999999999999999999999999999999999E+6000",
     "02:30 Asia/Tokyo", "2024-11-03 01:30 America/New_York"),  # ambiguous: Firebird takes the first 01:30
]
EXPECTED = {  # Python values the driver must return for ROWS
    1: (170141183460469231731687303715884105727, D("-1234567890123456789012345678901234.5678"),
        12345678901234567890123456789012345678, D("1.234567890123456E+300"),
        D("-1.234567890123456789012345678901234E-6000"),
        dt.time(10, 11, 12, 345600, tzinfo=NY), dt.datetime(2024, 7, 1, 10, 11, 12, 345600, tzinfo=NY)),
    2: (-(2**127), D("-0.0001"), -98765432109876543210, D("-0"), D("NaN"),
        dt.time(23, 59, 59, tzinfo=BUC), dt.datetime(2024, 1, 15, 23, 59, 59, tzinfo=BUC)),
    3: (2**64, D("-1.5000"), -1, D("Infinity"), D("-Infinity"),
        dt.time(1, 0, tzinfo=P0530), dt.datetime(2024, 3, 10, 1, 0, tzinfo=P0530)),
    4: (-(2**64) - 1, D("12.3400"), 2**63, D("1.000000000000000E+384"), D("1E-6143"),
        dt.time(0, 0, tzinfo=M0300), dt.datetime(1, 1, 1, 0, 0, tzinfo=M0300)),
}
COLS = ["i", "n", "d", "f16", "f34", "tt", "ts"]


def fb4_or_skip(conn):
    cur = conn.cursor()
    cur.execute("SELECT RDB$GET_CONTEXT('SYSTEM', 'ENGINE_VERSION') FROM RDB$DATABASE")
    if int(cur.fetchone()[0].split(".")[0]) < 4:
        pytest.skip("INT128, DECFLOAT and time zones need Firebird 4+")


@pytest.fixture
def conn():
    c = firebirdsql.create_database(host=HOST, port=PORT, user=USER, password=PASSWORD,
                                    database=f"{DB_DIR}/fbadmin_fb4types_{uuid.uuid4().hex[:8]}.fdb")
    try:
        fb4_or_skip(c)
        cur = c.cursor()
        cur.execute(TABLE)
        c.commit()
        for r in ROWS:
            cur.execute(f"insert into t values ({r[0]}, {r[1]}, {r[2]}, {r[3]}, '{r[4]}', '{r[5]}', "
                        f"time '{r[6]}', timestamp '{r[7]}')")
        c.commit()
        yield c
    finally:
        c.drop_database()


def same(a, b):
    if isinstance(a, D) and isinstance(b, D) and a.is_nan():
        return b.is_nan()
    return a == b and type(a) is type(b) and str(a) == str(b)


# ------------------------------------------------------------------ driver level
@requires_server
def test_decode_matches_server_text(conn):
    cur = conn.cursor()
    cur.execute("select id, " + ", ".join(COLS) + ", "
                + ", ".join(f"cast({c} as varchar(100))" for c in COLS) + " from t order by id")
    for row in cur.fetchall():
        values, texts = row[1:8], row[8:]
        if row[0] in EXPECTED:
            for got, want in zip(values, EXPECTED[row[0]], strict=True):
                assert same(got, want), (row[0], got, want)
        for v, text in zip(values, texts, strict=True):
            assert to_json(v) == text or D(str(to_json(v))) == D(text), (row[0], v, text)
    # the zone of time zone values survives: name or offset, as the server writes it
    cur.execute("select ts, tt from t where id = 6")
    ts, tt = cur.fetchone()
    assert fb_temporal_text(ts) == "2024-11-03 01:30:00.0000 America/New_York"
    assert fb_temporal_text(tt) == "02:30:00.0000 Asia/Tokyo"
    # 01:30 happens twice in New York that day: fold tells them apart, and they go back unchanged
    cur.execute("select timestamp '2024-11-03 05:30 UTC' at time zone 'America/New_York', "
                "timestamp '2024-11-03 06:30 UTC' at time zone 'America/New_York' from rdb$database")
    first, second = cur.fetchone()
    assert (first.fold, first.utcoffset(), second.fold, second.utcoffset()) == (
        0, dt.timedelta(hours=-4), 1, dt.timedelta(hours=-5))
    cur.execute("select cast(cast(? as timestamp with time zone) at time zone 'UTC' as varchar(50)), "
                "cast(cast(? as timestamp with time zone) at time zone 'UTC' as varchar(50)) from rdb$database",
                [first, second])
    assert cur.fetchone() == ("2024-11-03 05:30:00.0000 UTC", "2024-11-03 06:30:00.0000 UTC")


@requires_server
def test_extended_bind_and_regions(conn):
    """SET BIND ... TO EXTENDED (sqltypes 32748/32750) and a few zones whose Python names differ."""
    cur = conn.cursor()
    cur.execute("set bind of time zone to extended")
    cur.execute("select ts, tt from t where id = 1")
    assert cur.fetchone() == EXPECTED[1][6:4:-1]
    cur.execute("select cast(timestamp '2024-01-15 12:00 UTC' at time zone 'AET' as varchar(60)), "
                "timestamp '2024-01-15 12:00 UTC' at time zone 'AET', "
                "cast(timestamp '1890-06-01 00:00 UTC' at time zone 'America/New_York' as varchar(60)), "
                "timestamp '1890-06-01 00:00 UTC' at time zone 'America/New_York' from rdb$database")
    t1, v1, t2, v2 = cur.fetchone()
    assert fb_temporal_text(v1) == t1 == "2024-01-15 23:00:00.0000 AET"
    assert fb_temporal_text(v2) == t2  # local mean time: Firebird truncates the offset to minutes


@requires_server
def test_parameters(conn):
    cur = conn.cursor()
    for id_, expected in EXPECTED.items():
        for col, v in zip(COLS, expected, strict=True):
            if isinstance(v, D) and v.is_nan():
                continue  # NaN = NaN raises "Decimal float invalid operation" in Firebird
            # comparing with the NaN row raises too: DECFLOAT(34) only within this row
            where = f"id = {id_} and {col} = ?" if col == "f34" else f"{col} = ?"
            cur.execute(f"select id from t where {where}", [v])
            assert [r[0] for r in cur.fetchall()] == [id_], (col, v)
    # write them back (INSERT / UPDATE) and read them again
    for id_, values in EXPECTED.items():
        cur.execute("insert into t (id, " + ", ".join(COLS) + ") values (?, ?, ?, ?, ?, ?, ?, ?)",
                    [100 + id_, *values])
        cur.execute("select " + ", ".join(COLS) + " from t where id = ?", [100 + id_])
        for got, want in zip(cur.fetchone(), values, strict=True):
            assert same(got, want), (id_, got, want)
        cur.execute("update t set i = ?, n = ? where id = ?", [-(2**100), D("-0.5"), 100 + id_])
        cur.execute("select i, n from t where id = ?", [100 + id_])
        assert cur.fetchone() == (-(2**100), D("-0.5000"))
    conn.rollback()
    # Python values from elsewhere: other offsets, a region time, floats out of the int64 range
    cur.execute("select cast(? as timestamp with time zone), cast(? as time with time zone), "
                "cast(? as double precision), cast(? as decfloat(34)) from rdb$database",
                [dt.datetime(2024, 7, 1, 12, 0, tzinfo=dt.timezone.utc), dt.time(23, 59, 59, tzinfo=BUC),
                 1e300, D("-Infinity")])
    ts, tt, f, df = cur.fetchone()
    assert fb_temporal_text(ts) == "2024-07-01 12:00:00.0000 +00:00"
    assert fb_temporal_text(tt) == "23:59:59.0000 Europe/Bucharest"
    assert f == 1e300 and df == D("-Infinity")


def test_offsets_and_out_of_range_units():
    # offset zone ids are 1439 + minutes; region ids count down from 65535 (GMT)
    assert fb4types.zone_by_id(1439 + 330) == P0530
    assert fb4types.zone_by_id(65535).key == "GMT"
    assert str(fb4types.zone_by_id(65534)) == "ACT"  # not in the IANA database: an ICU alias
    raw = (-678575).to_bytes(4, "big", signed=True) + (36000000).to_bytes(4, "big")  # 0001-01-01 01:00 UTC
    raw += (1439 - 180).to_bytes(4, "big") + (-180).to_bytes(4, "big", signed=True)
    assert fb4types.decode(fb4types.SQL_TIMESTAMP_TZ, 0, raw) == "0000-12-31 22:00:00.0000 -03:00"
    assert fb4types.decode(fb4types.SQL_INT128, -2, (-1).to_bytes(16, "big", signed=True)) == D("-0.01")


# ------------------------------------------------------------------ API level
def base(sid, dbid):
    return f"/api/servers/{sid}/databases/{dbid}"


def run(client, sid, dbid, sql):
    r = client.post(base(sid, dbid) + "/execute", json={"sql": sql}).json()
    assert r["error"] is None, r["error"]
    return r


@pytest.fixture
def fb4db(client, server):
    created = []

    def make():
        path = f"{DB_DIR}/fbadmin_fb4api_{uuid.uuid4().hex[:8]}.fdb"
        r = client.post(f"/api/servers/{server['id']}/databases/create", json={"name": path.rsplit("/", 1)[1],
                                                                               "path": path})
        assert r.status_code == 201, r.text
        created.append(r.json()["id"])
        return r.json()["id"]

    sid = server["id"]
    a = make()
    major = int(run(client, sid, a, "SELECT RDB$GET_CONTEXT('SYSTEM', 'ENGINE_VERSION') FROM RDB$DATABASE")
                ["results"][0]["rows"][0][0].split(".")[0])
    try:
        if major < 4:
            pytest.skip("INT128, DECFLOAT and time zones need Firebird 4+")
        run(client, sid, a, TABLE + ";\n" + "\n".join(
            f"insert into t values ({r[0]}, {r[1]}, {r[2]}, {r[3]}, '{r[4]}', '{r[5]}', time '{r[6]}', "
            f"timestamp '{r[7]}');" for r in ROWS))
        yield sid, a, make
    finally:
        for dbid in created:
            client.post(f"/api/servers/{sid}/databases/{dbid}/drop")


def server_texts(client, sid, dbid):
    r = run(client, sid, dbid, "select id, " + ", ".join(f"cast({c} as varchar(100))" for c in COLS)
            + " from t order by id")
    return {row[0]: row[1:] for row in r["results"][0]["rows"]}


@requires_server
def test_api_grid_query_exports(client, fb4db):
    sid, dbid, make = fb4db
    b = base(sid, dbid)
    texts = server_texts(client, sid, dbid)

    def check(rows):
        for row in rows:
            for got, text in zip(row[1:], texts[row[0]], strict=True):
                assert got == text or D(got) == D(text), (row[0], got, text)  # big numbers as strings

    data = client.get(b + "/tables/T/data").json()
    assert [c["type"] for c in data["columns"]][1:] == [
        "INT128", "NUMERIC(*,4)", "INT128", "DECFLOAT(16)", "DECFLOAT(34)", "TIME WITH TIME ZONE",
        "TIMESTAMP WITH TIME ZONE"]
    check(data["rows"])
    assert data["rows"][0][1] == "170141183460469231731687303715884105727"
    assert data["rows"][2][7] == "2024-03-10 01:00:00.0000 +05:30"

    s = client.post("/api/query/sessions", json={"server_id": sid, "database_id": dbid}).json()["session_id"]
    r = client.post(f"/api/query/sessions/{s}/execute", json={"sql": "select * from t order by id"}).json()
    assert r["error"] is None
    check(r["results"][0]["rows"])
    client.delete(f"/api/query/sessions/{s}")

    # grid edits: change a time zone cell and an INT128 cell, keyed by values read from the grid
    row = data["rows"][0]
    upd = {"key": {"ID": 1}, "values": {"TS": "2025-02-03 04:05:06.7891 Europe/Bucharest",
                                        "I": "-170141183460469231731687303715884105728", "N": "-0.0001"}}
    assert client.post(b + "/tables/T/data", json={"updates": [upd]}).status_code == 200
    got = client.get(b + "/tables/T/data").json()["rows"][0]
    assert got[1] == "-170141183460469231731687303715884105728" and got[2] == "-0.0001"
    assert got[7] == "2025-02-03 04:05:06.7891 Europe/Bucharest"
    # unchanged values sent back as they were read round-trip exactly
    upd = {"key": {"ID": 1}, "values": {"TS": row[7], "TT": got[6], "I": row[1]}}
    assert client.post(b + "/tables/T/data", json={"updates": [upd]}).json()["affected"] == 1
    assert client.get(b + "/tables/T/data").json()["rows"][0][7] == row[7]
    # a table keyed by a time zone value
    run(client, sid, dbid, "create table k (ts timestamp with time zone not null primary key, v int);\n"
        "insert into k values (timestamp '2024-11-03 01:30 America/New_York', 1);")
    key = client.get(b + "/tables/K/data").json()["rows"][0][0]
    assert key == "2024-11-03 01:30:00.0000 America/New_York"
    r = client.post(b + "/tables/K/data", json={"updates": [{"key": {"TS": key}, "values": {"V": 2}}]})
    assert r.json()["affected"] == 1, r.text
    # plain TIMESTAMP keys/values travel as ISO 8601 ("T", 6 digits), which Firebird cannot parse
    run(client, sid, dbid, "create table p (t timestamp not null primary key, v timestamp);\n"
        "insert into p values (timestamp '2024-07-01 10:11:12.3456', null);")
    key = client.get(b + "/tables/P/data").json()["rows"][0][0]
    assert key == "2024-07-01T10:11:12.345600"
    r = client.post(b + "/tables/P/data", json={"updates": [{"key": {"T": key}, "values": {"V": key}}]})
    assert r.json()["affected"] == 1, r.text
    assert client.get(b + "/tables/P/data").json()["rows"][0] == [key, key]

    # exports
    texts = server_texts(client, sid, dbid)
    e = b + "/schema/export"
    js = client.get(e, params={"table": "T", "format": "json"}).json()
    for row in js:
        for c, text in zip(COLS, texts[row["ID"]], strict=True):
            assert row[c.upper()] == text or D(str(row[c.upper()])) == D(text)
    csv_rows = client.get(e, params={"table": "T", "format": "csv"}).text.splitlines()[1:]
    assert csv_rows[2].split(",")[1:] == [
        "18446744073709551616", "-1.5000", "-1", "Infinity", "-Infinity", "01:00:00.0000 +05:30",
        "2024-03-10 01:00:00.0000 +05:30"]
    from openpyxl import load_workbook

    x = client.get(e, params={"table": "T", "format": "xlsx"})
    xrows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    assert xrows[1][1] == "170141183460469231731687303715884105727" and xrows[1][2] == -0.0001
    assert xrows[6][2] == "99999999999999999999999999999999.9999" and xrows[6][3] == -2
    assert xrows[3][7] == "2024-03-10 01:00:00.0000 +05:30"
    sql = client.post(e, json={"sql": "select * from t order by id", "format": "sql", "table_name": "T"}).text
    # the INSERT script recreates the same values
    b2 = make()
    sid2 = sid
    run(client, sid2, b2, TABLE + ";\n" + sql)
    assert server_texts(client, sid2, b2) == texts


@requires_server
def test_api_copy_data(client, fb4db):
    sid, dbid, make = fb4db
    target = make()
    texts = server_texts(client, sid, dbid)
    job = client.post("/api/schema/copy", json={"source": {"server_id": sid, "db_id": dbid},
                                                 "target": {"server_id": sid, "db_id": target},
                                                 "tables": ["T"]}).json()
    for _ in range(600):
        res = client.get(f"/api/schema/jobs/{job['id']}/result").json()
        if res["status"] != "running":
            break
        time.sleep(0.1)
    assert res["status"] == "finished", json.dumps(res)[:2000]
    assert server_texts(client, sid, target) == texts
