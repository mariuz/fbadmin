"""Milestone 0.6 schema tools: extract, diff, ERD, import/export, copy, search, generator, report."""

import io
import json
import time
import uuid
from pathlib import Path

import pytest

from fbadmin.sqlsplit import split_sql

from .conftest import DB_DIR, requires_server

HERE = Path(__file__).parent
COMPLEX = (HERE / "complex_schema.sql").read_text()
CHANGES = (HERE / "complex_schema_changes.sql").read_text()


def base(sid, dbid):
    return f"/api/servers/{sid}/databases/{dbid}"


def run_sql(client, sid, dbid, sql):
    r = client.post(base(sid, dbid) + "/execute", json={"sql": sql})
    assert r.status_code == 200, r.text
    err = r.json()["error"]
    assert err is None, err
    return r.json()


@pytest.fixture
def make_db(client, server):
    created = []

    def make(sql: str | None = None):
        path = f"{DB_DIR}/fbadmin_schema_{uuid.uuid4().hex[:8]}.fdb"
        r = client.post(f"/api/servers/{server['id']}/databases/create", json={"name": path.rsplit('/', 1)[1],
                                                                               "path": path})
        assert r.status_code == 201, r.text
        dbid = r.json()["id"]
        created.append(dbid)
        if sql:
            run_sql(client, server["id"], dbid, sql)
        return dbid

    yield make
    for dbid in created:
        client.post(f"/api/servers/{server['id']}/databases/{dbid}/disconnect")
        client.post(f"/api/servers/{server['id']}/databases/{dbid}/drop")


def wait_job(client, job):
    for _ in range(600):
        r = client.get(f"/api/schema/jobs/{job['id']}/result").json()
        if r["status"] != "running":
            return r
        time.sleep(0.1)
    raise AssertionError("job did not finish")


def extract(client, sid, dbid, **params):
    r = client.get(base(sid, dbid) + "/schema/extract", params={"header": False, **params})
    assert r.status_code == 200, r.text
    return r.json()["sql"]


def diff(client, source, target, **extra):
    r = client.post("/api/schema/diff", json={"source": source, "target": target, **extra})
    assert r.status_code == 200, r.text
    return r.json()


def test_split_collation_case_insensitive():
    stmts = split_sql("CREATE COLLATION C1 FOR UTF8 FROM UNICODE CASE INSENSITIVE;\nSELECT CASE WHEN 1=1 THEN 2 END "
                      "FROM RDB$DATABASE;\nCREATE TABLE T (A INT);")
    assert len(stmts) == 3


# ----------------------------------------------------------------- extract
@requires_server
@pytest.mark.parametrize("set_term", [True, False])
def test_extract_roundtrip_complex(client, server, make_db, set_term):
    sid = server["id"]
    a = make_db(COMPLEX)
    script = extract(client, sid, a, set_term=set_term)
    assert ("SET TERM ^ ;" in script) == set_term
    # dependency order: tables before FKs, stubs before views, views before real bodies, triggers at the end
    pos = {k: script.index(k) for k in ("CREATE TABLE EMP ", "ADD CONSTRAINT FK_EMP_DEPT", "CREATE PROCEDURE SP_TREE",
                                        "CREATE VIEW V_TREE", "ALTER PROCEDURE SP_TREE", "CREATE TRIGGER EMP_BI",
                                        "GRANT SELECT ON EMP TO ROLE R_APP", "COMMENT ON TABLE EMP")}
    assert list(pos.values()) == sorted(pos.values()), pos
    assert script.index("CREATE VIEW V_EMP ") < script.index("CREATE VIEW V_EMP_DEPT") < script.index(
        "CREATE VIEW V_TOP")
    b = make_db(script)
    again = extract(client, sid, b, set_term=set_term)
    assert again == script


@requires_server
def test_extract_download(client, server, make_db):
    a = make_db("CREATE TABLE T1 (ID INTEGER NOT NULL PRIMARY KEY);")
    r = client.get(base(server["id"], a) + "/schema/extract", params={"download": True})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert "CREATE TABLE T1" in r.text and "SET SQL DIALECT 3;" in r.text


# -------------------------------------------------------------------- diff
@requires_server
def test_diff_migrate_both_directions(client, server, make_db):
    sid = server["id"]
    a = make_db(COMPLEX)
    b = make_db(COMPLEX + "\n" + CHANGES)

    res = diff(client, {"server_id": sid, "db_id": a}, {"server_id": sid, "db_id": b})
    kinds = {(i["type"], i["name"]): i["status"] for i in res["items"]}
    assert kinds[("tables", "EXTRA_T")] == "drop"
    assert kinds[("tables", "TMP_CALC")] == "create"
    assert kinds[("views", "V_EMP")] == "alter"
    assert kinds[("procedures", "SP_RAISE")] == "create"
    assert all(i["source_ddl"] or i["target_ddl"] for i in res["items"])
    run_sql(client, sid, b, res["script"])
    assert diff(client, {"server_id": sid, "db_id": a}, {"server_id": sid, "db_id": b})["items"] == []

    # the other way round: make the original schema look like the changed one
    c = make_db(COMPLEX + "\n" + CHANGES)
    res = diff(client, {"server_id": sid, "db_id": c}, {"server_id": sid, "db_id": a}, set_term=False)
    assert "SET TERM" not in res["script"]
    run_sql(client, sid, a, res["script"])
    after = diff(client, {"server_id": sid, "db_id": c}, {"server_id": sid, "db_id": a})
    assert after["items"] == [], [(i["type"], i["name"]) for i in after["items"]]
    assert "identical" in after["script"]


@requires_server
def test_diff_against_ddl_script(client, server, make_db):
    sid = server["id"]
    empty = make_db()
    res = diff(client, {"server_id": sid, "sql": COMPLEX}, {"server_id": sid, "db_id": empty})
    assert res["summary"]["create"] > 20 and res["source_name"] == "DDL script"
    run_sql(client, sid, empty, res["script"])
    assert diff(client, {"server_id": sid, "sql": COMPLEX}, {"server_id": sid, "db_id": empty})["items"] == []
    r = client.post("/api/schema/diff", json={"source": {"server_id": sid, "sql": "CREATE TABLE BROKEN ("},
                                              "target": {"server_id": sid, "db_id": empty}})
    assert r.status_code == 400 and "DDL script failed" in r.json()["detail"]


# --------------------------------------------------------------------- ERD
@requires_server
def test_erd(client, server, make_db):
    a = make_db(COMPLEX)
    d = client.get(base(server["id"], a) + "/schema/erd").json()
    names = {t["name"] for t in d["tables"]}
    assert {"EMP", "DEPT", "PROJECT", "EMP_PROJECT"} <= names
    rel = {(r["from_table"], r["to_table"]) for r in d["relations"]}
    assert ("EMP", "DEPT") in rel and ("DEPT", "DEPT") in rel and ("EMP_PROJECT", "PROJECT") in rel
    emp = next(t for t in d["tables"] if t["name"] == "EMP")
    cols = {c["name"]: c for c in emp["columns"]}
    assert cols["ID"]["pk"] and cols["DEPT_ID"]["fk"] and cols["BONUS"]["computed"]


# ------------------------------------------------------------ import/export
def upload(client, name, data: bytes):
    r = client.post("/api/schema/uploads", params={"filename": name}, content=data)
    assert r.status_code == 201, r.text
    return r.json()


@requires_server
def test_import_csv_create_and_errors(client, server, make_db):
    sid = server["id"]
    a = make_db()
    csv_data = "id;name;price;born;active\n1;Ann;10.50;2001-02-03;true\n2;Bob;7;1999-12-31;false\n3;\"Smith; Jr\";;;\n"
    up = upload(client, "people.csv", csv_data.encode())
    pv = client.post(f"/api/schema/uploads/{up['upload_id']}/preview", json={}).json()
    assert pv["delimiter"] == ";" and [c["name"] for c in pv["columns"]] == ["id", "name", "price", "born", "active"]
    types = {c["name"]: c["type"] for c in pv["columns"]}
    assert types["id"] == "INTEGER" and types["price"].startswith("NUMERIC") and types["born"] == "DATE"
    assert types["active"] == "BOOLEAN" and pv["rows"][2][1] == "Smith; Jr"
    cols = [{"source": c["name"], "target": c["suggested_name"], "type": c["type"]} for c in pv["columns"]]
    job = client.post(base(sid, a) + "/schema/import", json={"upload_id": up["upload_id"], "table": "PEOPLE",
                                                             "create": True, "columns": cols}).json()
    res = wait_job(client, job)
    assert res["status"] == "finished" and res["inserted"] == 3, res
    rows = run_sql(client, sid, a, "SELECT ID, NAME, PRICE, BORN, ACTIVE FROM PEOPLE ORDER BY ID")["results"][0]["rows"]
    assert rows[0] == [1, "Ann", "10.50", "2001-02-03", True] and rows[2][2] is None

    # a bad row: abort rolls everything back, skip imports the rest and reports the row
    bad = b"id,name\n10,ok\nxx,bad\n11,ok2\n"
    mapping = [{"source": "id", "target": "ID"}, {"source": "name", "target": "NAME"}]
    up = upload(client, "bad.csv", bad)
    res = wait_job(client, client.post(base(sid, a) + "/schema/import", json={
        "upload_id": up["upload_id"], "table": "PEOPLE", "columns": mapping, "delete_upload": False}).json())
    assert res["status"] == "failed" and "Row 3" in res["error"]
    assert run_sql(client, sid, a, "SELECT COUNT(*) FROM PEOPLE")["results"][0]["rows"][0][0] == 3
    res = wait_job(client, client.post(base(sid, a) + "/schema/import", json={
        "upload_id": up["upload_id"], "table": "PEOPLE", "columns": mapping, "on_error": "skip"}).json())
    assert res["status"] == "finished" and res["inserted"] == 2 and res["failed"] == 1
    assert res["errors"][0]["row"] == 3
    assert client.post(f"/api/schema/uploads/{up['upload_id']}/preview", json={}).status_code == 404


@requires_server
def test_import_json_and_xlsx(client, server, make_db):
    from openpyxl import Workbook, load_workbook

    sid = server["id"]
    a = make_db("CREATE TABLE ITEMS (ID INTEGER NOT NULL PRIMARY KEY, LABEL VARCHAR(20), QTY NUMERIC(10,2), "
                "TS TIMESTAMP, META BLOB SUB_TYPE TEXT);")
    data = [{"id": i, "label": f"item {i}", "qty": i * 1.5, "ts": "2024-05-06T07:08:09", "meta": {"k": i}}
            for i in range(1, 501)]
    up = upload(client, "items.json", json.dumps(data).encode())
    pv = client.post(f"/api/schema/uploads/{up['upload_id']}/preview", json={}).json()
    assert [c["name"] for c in pv["columns"]] == ["id", "label", "qty", "ts", "meta"]
    mapping = [{"source": s, "target": t} for s, t in
               (("id", "ID"), ("label", "LABEL"), ("qty", "QTY"), ("ts", "TS"), ("meta", "META"))]
    res = wait_job(client, client.post(base(sid, a) + "/schema/import", json={
        "upload_id": up["upload_id"], "table": "ITEMS", "columns": mapping}).json())
    assert res["status"] == "finished" and res["inserted"] == 500, res
    r = run_sql(client, sid, a, "SELECT QTY, TS, META FROM ITEMS WHERE ID = 3")["results"][0]["rows"][0]
    assert r[0] == "4.50" and r[1].startswith("2024-05-06T07:08:09") and json.loads(r[2]) == {"k": 3}

    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "LABEL", "QTY"])
    for i in range(1000, 1010):
        ws.append([i, f"x{i}", i / 4])
    buf = io.BytesIO()
    wb.save(buf)
    up = upload(client, "items.xlsx", buf.getvalue())
    pv = client.post(f"/api/schema/uploads/{up['upload_id']}/preview", json={}).json()
    assert pv["format"] == "xlsx" and len(pv["rows"]) == 10 and pv["sheets"]
    mapping = [{"source": c, "target": c} for c in ("ID", "LABEL", "QTY")]
    res = wait_job(client, client.post(base(sid, a) + "/schema/import", json={
        "upload_id": up["upload_id"], "table": "ITEMS", "columns": mapping}).json())
    assert res["inserted"] == 10, res
    assert run_sql(client, sid, a, "SELECT COUNT(*) FROM ITEMS")["results"][0]["rows"][0][0] == 510

    # exports
    b = base(sid, a) + "/schema/export"
    csv_text = client.get(b, params={"table": "ITEMS", "format": "csv"}).text
    assert csv_text.splitlines()[0] == "ID,LABEL,QTY,TS,META" and len(csv_text.splitlines()) == 511
    js = client.get(b, params={"table": "ITEMS", "format": "json"}).json()
    assert len(js) == 510 and js[0]["LABEL"] == "item 1"
    sql = client.post(b, json={"sql": "SELECT ID, LABEL FROM ITEMS WHERE ID < 3 ORDER BY ID", "format": "sql",
                               "table_name": "COPY"}).text
    assert sql == "INSERT INTO COPY (ID, LABEL) VALUES (1, 'item 1');\nINSERT INTO COPY (ID, LABEL) VALUES (2, 'item 2');\n"
    x = client.get(b, params={"table": "ITEMS", "format": "xlsx"})
    assert x.status_code == 200
    rows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("ID", "LABEL", "QTY") and len(rows) == 511
    assert client.post(b, json={"sql": "DELETE FROM ITEMS", "format": "csv"}).status_code == 400


# --------------------------------------------------------- copy / generate
@requires_server
def test_generator_and_copy(client, server, make_db):
    sid = server["id"]
    a = make_db(COMPLEX)
    gen = base(sid, a) + "/schema/generate"
    for table, n in (("DEPT", 15), ("EMP", 60), ("PROJECT", 20), ("EMP_PROJECT", 40), ("NO_PK", 25)):
        res = wait_job(client, client.post(gen, json={"table": table, "rows": n, "seed": 7}).json())
        assert res["status"] == "finished" and res["inserted"] == n and res["failed"] == 0, (table, res)
    q = lambda s: run_sql(client, sid, a, s)["results"][0]["rows"]  # noqa: E731
    assert q("SELECT COUNT(*) FROM EMP WHERE STATUS NOT IN ('A', 'I') OR SALARY < 0 OR EMAIL NOT LIKE '%@%'") == [[0]]
    assert q("SELECT COUNT(*) FROM EMP e WHERE e.DEPT_ID IS NOT NULL AND NOT EXISTS "
             "(SELECT 1 FROM DEPT d WHERE d.ID = e.DEPT_ID)") == [[0]]
    assert q("SELECT COUNT(DISTINCT CODE) FROM PROJECT") == [[20]]
    # overrides: fixed value, list, null
    res = wait_job(client, client.post(gen, json={"table": "NO_PK", "rows": 5, "overrides": {
        "A": {"mode": "value", "value": 42}, "B": {"mode": "list", "values": "x, y"}}}).json())
    assert res["inserted"] == 5
    assert q("SELECT COUNT(*) FROM NO_PK WHERE A = 42 AND B IN ('x', 'y')") == [[5]]

    # copy everything into an empty database (tables are created)
    b = make_db()
    tables = ["DEPT", "EMP", "PROJECT", "EMP_PROJECT", "NO_PK"]
    job = client.post("/api/schema/copy", json={"source": {"server_id": sid, "db_id": a},
                                                "target": {"server_id": sid, "db_id": b},
                                                "tables": tables, "batch_size": 7}).json()
    res = wait_job(client, job)
    assert res["status"] == "finished", res
    assert res["tables"] == {"DEPT": 15, "EMP": 60, "PROJECT": 20, "EMP_PROJECT": 40, "NO_PK": 30}
    log = client.get(f"/api/jobs/{job['id']}").json()["log"]
    assert any("EMP: 7 / 60 rows" in line for line in log)
    q2 = lambda s: run_sql(client, sid, b, s)["results"][0]["rows"]  # noqa: E731
    assert q2("SELECT COUNT(*), SUM(ID) FROM EMP") == q("SELECT COUNT(*), SUM(ID) FROM EMP")
    fks = client.get(base(sid, b) + "/objects/tables/EMP").json()["constraints"]
    assert any(c["type"] == "FOREIGN KEY" and c["ref_table"] == "DEPT" for c in fks)
    # identity restarted after the copied values
    run_sql(client, sid, b, "INSERT INTO DEPT (NAME) VALUES ('new one')")
    # copying again into existing tables with truncate
    res = wait_job(client, client.post("/api/schema/copy", json={
        "source": {"server_id": sid, "db_id": a}, "target": {"server_id": sid, "db_id": b},
        "tables": ["NO_PK"], "truncate": True, "create_tables": False}).json())
    assert res["tables"] == {"NO_PK": 30}
    assert q2("SELECT COUNT(*) FROM NO_PK") == [[30]]


# ------------------------------------------------------------------ search
@requires_server
def test_search(client, server, make_db):
    a = make_db(COMPLEX)
    s = base(server["id"], a) + "/schema/search"
    items = client.get(s, params={"q": "sp_t"}).json()["items"]
    assert {(i["type"], i["name"]) for i in items} >= {("procedures", "SP_TREE"), ("procedures", "SP_TOP_COUNT")}
    items = client.get(s, params={"q": "department", "scopes": "comment"}).json()["items"]
    assert ("procedures", "SP_TREE") in {(i["type"], i["name"]) for i in items}
    assert all(i["match"] == "comment" for i in items)
    items = client.get(s, params={"q": "pk_hr.rate", "scopes": "source"}).json()["items"]
    hit = next(i for i in items if i["name"] == "F_BONUS")
    assert hit["match"] == "source" and "PK_HR.RATE()" in hit["snippet"] and hit["line"] == 2
    items = client.get(s, params={"q": "salary", "types": "columns", "scopes": "name"}).json()["items"]
    assert {(i["name"], i["parent"]) for i in items} == {("SALARY", "EMP")}


# ------------------------------------------------------------------ report
@requires_server
def test_report(client, server, make_db):
    sid = server["id"]
    a = make_db(COMPLEX)
    # changing a package header leaves its body invalid
    run_sql(client, sid, a, "SET TERM ^ ;\nCREATE OR ALTER PACKAGE PK_HR AS\nBEGIN\n  FUNCTION RATE RETURNS "
                            "NUMERIC(5,2);\n  PROCEDURE RAISE_ALL (PCT INTEGER);\n  FUNCTION X RETURNS INT;\nEND^\n"
                            "SET TERM ; ^")
    secs = {s["id"]: {i["name"] for i in s["items"]}
            for s in client.get(base(sid, a) + "/schema/report").json()["sections"]}
    assert secs["inactive_indices"] == {"IX_PROJECT_BUDGET"}
    assert secs["inactive_triggers"] == {"EMP_AU"}
    assert secs["unused_domains"] == {"D_UNUSED"}
    assert secs["unused_exceptions"] == {"E_UNUSED"}
    assert secs["unused_sequences"] == {"SEQ_UNUSED"}
    assert secs["tables_without_pk"] == {"NO_PK"}
    assert secs["invalid_packages"] == {"PK_HR"}
    assert "invalid_procedures" in secs and "invalid_triggers" in secs
